#!/usr/bin/env python3
"""SOP-inspired MIR cleanup. Analyse -> review Confirm/Reason/Unit -> apply.

No candidate is automatically removed. See MIR_CLEANUP_README.md.
The original modelling module is imported for its established IO/preprocessing only.
"""
from __future__ import annotations
import argparse
import sys
from datetime import datetime
from decimal import Decimal
import hashlib
import json
from pathlib import Path
from functools import lru_cache
import numpy as np
import pandas as pd
from scipy.spatial.distance import cdist
from sklearn.decomposition import PCA
from sklearn.cross_decomposition import PLSRegression
from threadpoolctl import threadpool_limits

ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = ROOT if (ROOT / 'Python code for MIR').is_dir() else ROOT.parent
DEFAULT_SOURCE = PROJECT_ROOT / 'Python code for MIR/data/reference/reference_value_ZL_trt.xlsx'
DEFAULT_OUTPUT = PROJECT_ROOT / 'Python code for MIR/output/MIR cleanup'

# EDIT HERE: False uses your list below; True analyses every property sheet.
CONFIG = {
    'concentration_limit': 'disable',  # 'enable' enforces the percentage below.
    'concentration_limit_percent': 1.0,  # Percent: 5.0 means 5%; valid range 0–100.
    'run_all_properties': False,
    'property_sheets': ['202_STC', '202_STN'],
}
# Command-line --properties / --all-properties overrides these settings.
COLS = ['Sample', 'Reference Value', 'File Name', 'Group']
REASONS = {'range', 'redundancy', 'spectral_quality', 'concentration', 'historical_unverified'}

@lru_cache(maxsize=1)
def backend():
    if str(PROJECT_ROOT) not in sys.path:
        sys.path.insert(0, str(PROJECT_ROOT))
    import soil_mir_plsr_Nested_LOGO_stepwise_regions_rank_reuse as module
    return module


def fingerprint(path: str | Path) -> str:
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def write_json(path: Path, data: dict) -> None:
    path.write_text(json.dumps(data, indent=2, allow_nan=False), encoding='utf-8')


def validate_frame(frame: pd.DataFrame) -> pd.DataFrame:
    if not set(COLS).issubset(frame.columns):
        raise ValueError(f'Required columns: {COLS}')
    frame = frame.dropna(how='all').copy()
    if frame.empty or frame[COLS].isna().any().any():
        raise ValueError('Empty dataset or missing required values.')
    for col in ['Sample', 'File Name', 'Group']:
        if frame[col].astype(str).str.strip().eq('').any():
            raise ValueError(f'Blank {col}')
    if frame['File Name'].duplicated().any():
        raise ValueError('Duplicate reference filenames.')
    if not np.isfinite(pd.to_numeric(frame['Reference Value'], errors='raise')).all():
        raise ValueError('Non-finite reference values.')
    if (frame.groupby('Sample')[['Reference Value', 'Group']].nunique() > 1).any().any():
        raise ValueError('Inconsistent reference value or treatment within a sample.')
    return frame.reset_index(drop=True)


def fit_pca(X: np.ndarray) -> tuple:
    X = np.asarray(X, dtype=float)
    if X.ndim != 2 or len(X) < 2 or not np.isfinite(X).all():
        raise ValueError('PCA requires at least two finite spectra.')
    if np.all(X == X[0]):
        raise ValueError('PCA has no nonzero components: spectra are constant.')
    model = PCA(n_components=min(15, len(X) - 1, X.shape[1]), svd_solver='full')
    scores = model.fit_transform(X)
    tolerance = np.finfo(float).eps * max(X.shape) * model.singular_values_[0]
    n = int(np.sum(model.singular_values_ > tolerance))
    if n == 0:
        raise ValueError('PCA has no nonzero components: spectra are constant.')
    return model, scores[:, :n], model.explained_variance_ratio_[:n]


def kennard_stone(scores: np.ndarray, n_select: int) -> np.ndarray:
    """Deterministic farthest-pair / maximin sampling; input order breaks ties."""
    n = len(scores)
    if not 1 <= n_select < n:
        raise ValueError('Subset size must be between 1 and n-1.')
    distances = cdist(scores, scores)
    a, b = np.unravel_index(np.argmax(distances), distances.shape)
    selected = [int(a)]
    if n_select > 1:
        selected.append(int(b if b != a else (a + 1) % n))
    while len(selected) < n_select:
        nearest = distances[:, selected].min(axis=1)
        nearest[selected] = -np.inf
        selected.append(int(np.argmax(nearest)))
    return np.array(selected)


def diagnose_half(X: np.ndarray, frame: pd.DataFrame, train: np.ndarray,
                  test: np.ndarray, cfg: dict, label: str,
                  fixed_selection: tuple[str, int] | None = None) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Select on the opposite half; diagnose calibration records only."""
    mod = backend()
    groups = frame['Sample'].astype(str).to_numpy()
    y = frame['Reference Value'].to_numpy(float)
    output = pd.DataFrame(index=frame.index)
    output['Diagnostic Note'] = ''
    searches = []
    if len(set(groups[train])) < 2 or len(set(groups[test])) < 2:
        raise ValueError('At least two retained samples are required in each half.')
    yt, lam = mod.apply_transform(y[train], cfg['transform'])
    candidates = []
    for prep in ([fixed_selection[0]] if fixed_selection else mod.PREPROCESSING_NAMES):
        fitted, xp = mod.fit_transform_preprocessor(prep, X[train], cfg)
        xv = mod.transform_preprocessor(fitted, X[test])
        rank = min(15, len(xp) - 2, xp.shape[1], np.linalg.matrix_rank(xp - xp.mean(axis=0)))
        if rank < 1 or np.var(yt) == 0:
            continue
        if fixed_selection and not 1 <= fixed_selection[1] <= rank:
            raise ValueError(f"Half {label} cannot support A's selected rank {fixed_selection[1]} (available rank {rank}).")
        for r in ([fixed_selection[1]] if fixed_selection else range(1, rank + 1)):
            model = PLSRegression(n_components=r, scale=False).fit(xp, yt)
            prediction = mod.back_transform(model.predict(xv).ravel(), cfg['transform'], lam)
            means = pd.DataFrame({'sample': groups[test], 'true': y[test], 'pred': prediction}).groupby('sample').mean()
            error = float(np.sqrt(np.mean((means.true - means.pred) ** 2)))
            if not np.isfinite(error):
                raise ValueError('Non-finite diagnostic predictions.')
            searches.append({'Half': label, 'Preprocessing': prep, 'Rank': r, 'Opposite-half RMSEP': error})
            candidates.append((error, r, prep, fitted, xp, rank))
    if not candidates:
        output.loc[train, 'Diagnostic Note'] = 'Concentration diagnostic unavailable: constant response or rank zero.'
        return output, pd.DataFrame(searches)
    best = min(candidates, key=lambda c: (c[0], c[1], c[2]))
    error, selected_rank, prep, fitted, xp, rank = best
    pls = PLSRegression(n_components=rank, scale=False).fit(xp, yt)
    pred_t = pls.predict(xp).ravel()
    residual = yt - pred_t
    t = pls.x_scores_
    leverage = np.sum((t @ np.linalg.pinv(t.T @ t)) * t, axis=1).clip(0, 1)
    mse = float(np.mean(residual ** 2))
    stud = residual / np.sqrt(mse * (1 - leverage + 1e-10)) if mse > 0 else np.zeros(len(xp))
    values = {'Half': label, 'Preprocessing': prep, 'Selection Rank': selected_rank,
              'Diagnostic Rank': rank, 'Opposite-half RMSEP': error,
              'Transformed Reference': yt, 'Transformed Prediction': pred_t,
              'Transformed Residual': residual, 'Studentised Residual': stud,
              'Leverage': leverage, 'Concentration Candidate': np.abs(stud) > 2.5,
              'Calibration Prediction': mod.back_transform(pred_t, cfg['transform'], lam)}
    for name, value in values.items():
        output.loc[train, name] = value
    return output.loc[train].copy(), pd.DataFrame(searches)


def concentration_diagnostics(X: np.ndarray, frame: pd.DataFrame, scores: np.ndarray,
                              cfg: dict) -> tuple[pd.DataFrame, pd.DataFrame]:
    mod = backend()
    groups = frame['Sample'].astype(str).to_numpy()
    y = frame['Reference Value'].to_numpy(float)
    avg = pd.DataFrame(scores).assign(Sample=groups).groupby('Sample', sort=True).mean()
    output = pd.DataFrame(index=frame.index)
    output['Diagnostic Note'] = ''
    searches = []
    if len(avg) < 4:
        output['Diagnostic Note'] = 'Concentration diagnostic unavailable: fewer than four samples.'
        output['Concentration Candidate'] = False
        return output, pd.DataFrame()
    chosen = kennard_stone(avg.to_numpy(), len(avg) // 2)
    first = np.isin(groups, avg.index[chosen])
    for label, train in [('A', first), ('B', ~first)]:
        test = ~train
        diagnostic, search = diagnose_half(X, frame, train, test, cfg, label)
        for name in diagnostic:
            output.loc[diagnostic.index, name] = diagnostic[name]
        searches.extend(search.to_dict(orient='records'))
    if 'Concentration Candidate' not in output:
        output['Concentration Candidate'] = False
    output['Concentration Candidate'] = output['Concentration Candidate'].eq(True)
    return output, pd.DataFrame(searches)


def analyse_property(frame: pd.DataFrame, spectra: dict, axis: np.ndarray, cfg: dict,
                     bounds: dict | None = None) -> tuple:
    frame = validate_frame(frame)
    missing = set(frame['File Name']) - set(spectra)
    if missing:
        raise ValueError(f'Missing spectra: {sorted(missing)[:10]}')
    axis = backend().validate_wavenumbers(axis)
    mask = (axis >= 600) & (axis <= 4000)
    if cfg.get('exclude_co2'):
        mask &= ~((axis >= 2300) & (axis <= 2400))
    X = np.array([spectra[f] for f in frame['File Name']], dtype=float)[:, mask]
    selected_axis = axis[mask]
    breaks = np.flatnonzero(np.abs(np.diff(selected_axis)) > 1.5 * np.median(np.abs(np.diff(axis)))) + 1
    cfg = dict(cfg, _segment_lengths=[len(s) for s in np.split(selected_axis, breaks)])
    model, scores, variance = fit_pca(X)
    distance = cdist(scores, scores)
    np.fill_diagonal(distance, np.inf)
    nearest = distance.argmin(axis=1)
    diag, search = concentration_diagnostics(X, frame, scores, cfg)
    report = pd.concat([frame, diag], axis=1)
    report['Nearest File'] = frame['File Name'].to_numpy()[nearest]
    report['Nearest Distance'] = distance[np.arange(len(frame)), nearest]
    report['Same Sample Neighbour'] = frame['Sample'].to_numpy() == frame['Sample'].to_numpy()[nearest]
    report['Similarity Order'] = report['Nearest Distance'].rank(method='min').astype(int)
    report['Range Candidate'] = False
    bounds = bounds or {}
    lo, hi = bounds.get('min'), bounds.get('max')
    if any(v is not None and not np.isfinite(v) for v in [lo, hi]) or (lo is not None and hi is not None and lo > hi):
        raise ValueError('Invalid reference bounds.')
    if lo is not None:
        report['Range Candidate'] |= frame['Reference Value'] < lo
    if hi is not None:
        report['Range Candidate'] |= frame['Reference Value'] > hi
    report['Confirm'] = ''
    report['Reason'] = np.where(report['Range Candidate'], 'range', np.where(report['Concentration Candidate'], 'concentration', ''))
    report['Unit'] = np.where(report['Reason'].ne(''), 'sample', 'spectrum')
    report['Comment'] = ''
    records = []
    for i, row in report.iterrows():
        status = 'range candidate' if row['Range Candidate'] else 'concentration candidate' if row['Concentration Candidate'] else 'unconfirmed / retained'
        records.append({'sample': str(row['Sample']), 'file': str(row['File Name']), 'group': str(row['Group']),
                        'value': float(row['Reference Value']), 'scores': scores[i].tolist(), 'status': status,
                        'excluded': False, 'nearest': str(row['Nearest File']), 'distance': float(row['Nearest Distance'])})
    payload = {'variance': variance.tolist(), 'records': records, 'basis': 'PCA fitted before exclusions in this review pass'}
    return report, search, payload, X, model


def confirmed_rows(review: pd.DataFrame) -> pd.DataFrame:
    required = {'Property', 'Sample', 'File Name', 'Confirm', 'Reason', 'Unit'}
    if not required.issubset(review):
        raise ValueError(f'Review is missing {required - set(review.columns)}')
    values = review['Confirm'].fillna('').astype(str).str.strip().str.lower()
    if not values.isin(['', 'yes', 'no']).all():
        raise ValueError('Confirm must be blank, YES, or NO.')
    result = review.loc[values.eq('yes')].copy()
    for col in ['Reason', 'Unit']:
        result[col] = result[col].fillna('').astype(str).str.strip().str.lower()
    if not result['Reason'].isin(REASONS).all():
        raise ValueError(f'Confirmed rows need one of these reasons: {sorted(REASONS)}')
    if not result['Unit'].isin(['sample', 'spectrum']).all():
        raise ValueError('Unit must be sample or spectrum.')
    if ((result.Reason == 'concentration') & (result.Unit != 'sample')).any():
        raise ValueError('Concentration exclusions must use sample units.')
    return result


def concentration_limit_setting(value: str | None = None) -> str:
    value = CONFIG['concentration_limit'] if value is None else value
    if value not in ('enable', 'disable'):
        raise ValueError("concentration_limit must be 'enable' or 'disable'.")
    return value


def concentration_limit_percent_setting() -> Decimal:
    try:
        value = Decimal(str(CONFIG['concentration_limit_percent']))
    except (ValueError, ArithmeticError):
        raise ValueError('concentration_limit_percent must be a finite number from 0 to 100.') from None
    if not value.is_finite() or not 0 <= value <= 100:
        raise ValueError('concentration_limit_percent must be a finite number from 0 to 100.')
    return value


def apply_confirmations(frame: pd.DataFrame, review: pd.DataFrame, property_name: str,
                        historical_replay: bool = False,
                        concentration_limit: str | None = None) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Idempotent when evaluated against the same original source and review."""
    frame = validate_frame(frame)
    prop_review = review.loc[review.Property.eq(property_name)].copy()
    if prop_review['File Name'].duplicated().any():
        raise ValueError('Duplicate/conflicting review filenames.')
    indexed = frame.set_index('File Name')
    for _, r in prop_review.iterrows():
        if r['File Name'] not in indexed.index or str(indexed.loc[r['File Name'], 'Sample']) != str(r['Sample']):
            raise ValueError(f'Unknown filename or mismatched sample: {r["File Name"]}')
    yes = confirmed_rows(prop_review)
    if not historical_replay and yes.Reason.eq('historical_unverified').any():
        raise ValueError('Historical replay requires --historical-replay; otherwise assign an evidence-based reason.')
    no = set(prop_review.loc[prop_review.Confirm.fillna('').astype(str).str.strip().str.lower().eq('no'), 'File Name'])
    limit = concentration_limit_setting(concentration_limit)
    percent = concentration_limit_percent_setting()
    cap = int(Decimal(frame.Sample.nunique()) * percent / Decimal(100))
    if limit == 'enable' and yes.loc[yes.Reason.eq('concentration'), 'Sample'].nunique() > cap:
        raise ValueError(f'{property_name}: concentration exclusions exceed combined allowance of {cap} samples ({percent}% limit).')
    reasons: dict[str, set] = {}
    for _, row in yes.iterrows():
        targets = frame.loc[frame.Sample.eq(row.Sample), 'File Name'] if row.Unit == 'sample' else [row['File Name']]
        for f in targets:
            if f in no:
                raise ValueError(f'Conflicting NO decision and sample exclusion: {f}')
            reasons.setdefault(f, set()).add(row.Reason)
    if len(reasons) == len(frame):
        raise ValueError('Cannot exclude the entire dataset.')
    audit = frame.loc[frame['File Name'].isin(reasons)].copy()
    audit['Property'] = property_name
    audit['Reasons'] = [', '.join(sorted(reasons[f])) for f in audit['File Name']]
    return frame.loc[~frame['File Name'].isin(reasons)].copy(), audit


def historical_comparison(source: dict, properties: list[str], excluded: dict | None = None) -> pd.DataFrame:
    rows = []
    for prop in properties:
        if prop + '_OPUS' not in source:
            continue
        original = validate_frame(source[prop])
        historical = validate_frame(source[prop + '_OPUS'])
        if not set(historical['File Name']).issubset(original['File Name']):
            raise ValueError('Historical sheet contains unknown filenames.')
        for _, row in original.iterrows():
            old = row['File Name'] not in set(historical['File Name'])
            new = row['File Name'] in (excluded or {}).get(prop, set())
            rows.append({'Property': prop, 'Sample': row.Sample, 'File Name': row['File Name'],
                         'Historical Excluded': old, 'Python Excluded': new,
                         'Comparison': 'both' if old and new else 'historical only' if old else 'Python only' if new else 'neither',
                         'Historical Reason': 'unverified' if old else ''})
    return pd.DataFrame(rows)


def export_figures(payload: dict, output: Path) -> None:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.backends.backend_pdf import PdfPages
    with PdfPages(output / 'PCA_report.pdf') as pdf:
        for prop, data in payload.items():
            records = data['records']; scores = np.array([r['scores'] for r in records])
            excluded = np.array([r['excluded'] for r in records])
            candidate = np.array(['candidate' in r['status'] for r in records])
            fig, axes = plt.subplots(1, 3, figsize=(16, 4.5))
            ypc = min(1, scores.shape[1]-1)
            for keep, name, color in [(~excluded & ~candidate, 'Retained', '#237a95'), (~excluded & candidate, 'Candidate', '#e29422'), (excluded, 'Confirmed excluded', '#bb354d')]:
                axes[0].scatter(scores[keep, 0], scores[keep, ypc], s=12, label=name, color=color)
            axes[0].set(xlabel=f'PC1 ({data["variance"][0]:.1%})', ylabel=f'PC{ypc+1} ({data["variance"][ypc]:.1%})')
            axes[0].legend(fontsize=8)
            axes[1].bar(np.arange(1, len(data['variance'])+1), np.array(data['variance'])*100)
            axes[1].set(xlabel='Principal component', ylabel='Explained variance (%)')
            unique = {r['sample']: r['value'] for r in records}
            axes[2].hist(list(unique.values()), bins=20, color='#237a95')
            axes[2].set(xlabel='Reference value (original units)', ylabel='Samples')
            fig.suptitle(f'{prop} — {data["basis"]}')
            fig.tight_layout(); pdf.savefig(fig); fig.savefig(output / f'{prop}_overview.png', dpi=180); plt.close(fig)


def render_viewer(payload: dict, path: Path) -> None:
    from plotly.offline import get_plotlyjs
    template = (ROOT / 'mir_cleanup_viewer.html').read_text(encoding='utf-8')
    data = json.dumps(payload, allow_nan=False).replace('<', '\\u003c')
    path.write_text(template.replace('/*PLOTLY_BUNDLE*/', get_plotlyjs()).replace('/*PCA_DATA*/{}', data), encoding='utf-8')


def select_properties(source: dict[str, pd.DataFrame], requested: list[str] | str | None = None) -> list[str]:
    """Resolve explicit sheet names or all data sheets, preserving workbook order."""
    if requested is None:
        requested = 'all' if CONFIG['run_all_properties'] else CONFIG['property_sheets']
    names = [requested] if isinstance(requested, str) else list(requested)
    all_mode = len(names) == 1 and str(names[0]).lower() in {'all', 'auto'}
    if all_mode:
        names = [name for name, frame in source.items() if set(COLS).issubset(frame.columns)]
    elif any(str(name).lower() in {'all', 'auto'} for name in names):
        raise ValueError('Use all by itself, not alongside individual sheet names.')
    names = list(dict.fromkeys(names))
    if not names:
        raise ValueError('No property sheets selected or discovered.')
    for name in names:
        if name not in source:
            raise ValueError(f'Unknown property sheet: {name}. Available sheets: {list(source)}')
        if not set(COLS).issubset(source[name].columns):
            raise ValueError(f'{name} is not a property sheet; requires {COLS}.')
    return names


def metadata_for_properties(metadata: dict, properties: list[str]) -> dict:
    """Historical OPUS subsets inherit their original property's metadata."""
    result = dict(metadata)
    for prop in properties:
        if prop not in result and prop.endswith('_OPUS') and prop[:-5] in result:
            result[prop] = dict(result[prop[:-5]])
    return result


def analyse(source_path: Path, output: Path, properties: list[str] | str | None = None, bounds: dict | None = None,
            prior: Path | None = None) -> None:
    source_path = source_path.resolve()
    source = pd.read_excel(source_path, sheet_name=None)
    requested = properties if properties is not None else ('all' if CONFIG['run_all_properties'] else CONFIG['property_sheets'])
    requested_names = [requested] if isinstance(requested, str) else list(requested)
    all_mode = len(requested_names) == 1 and str(requested_names[0]).lower() in {'all', 'auto'}
    properties = select_properties(source, requested)
    print(f'Properties selected ({len(properties)}): {", ".join(properties)}')
    original_frames = {}
    selection_issues = []
    for prop in properties:
        try:
            original_frames[prop] = validate_frame(source[prop])
        except ValueError as exc:
            if not all_mode:
                raise ValueError(f'{prop}: {exc}') from exc
            selection_issues.append({'Property': prop, 'Issue': str(exc)})
            print(f'Skipping invalid property {prop}: {exc}')
    properties = list(original_frames)
    if not properties:
        raise ValueError('No valid property datasets remain after validation.')
    mod = backend(); base = mod.CONFIG.copy()
    base.update(reference_excel=str(source_path), spectra_cache=None)
    with pd.ExcelFile(source_path) as book:
        metadata = metadata_for_properties(mod.load_property_metadata(base, book), properties)
    previous_audit = pd.DataFrame(); previous_payload = None; iteration = 1
    if prior:
        manifest = json.loads((prior / 'manifest.json').read_text(encoding='utf-8'))
        if manifest['source_sha256'] != fingerprint(source_path) or manifest['properties'] != properties:
            raise ValueError('Prior pass source/properties do not match.')
        iteration = manifest['pass'] + 1
        if iteration > 3:
            raise ValueError('At most three review passes are supported.')
        if not (prior / 'confirmed_audit.csv').exists():
            raise ValueError('Apply the prior review before starting another pass.')
        previous_audit = pd.read_csv(prior / 'confirmed_audit.csv')
        previous_payload = json.loads((prior / 'pca_data.json').read_text(encoding='utf-8'))
    # Validate metadata before creating outputs or loading the spectral library.
    for prop in properties:
        cfg = mod.build_property_config(base, prop, metadata)
        if cfg['metadata_warning']:
            raise ValueError(f'{prop}: {cfg["metadata_warning"]}')
    output.mkdir(parents=True, exist_ok=False)
    if selection_issues:
        pd.DataFrame(selection_issues).to_csv(output / 'skipped_properties.csv', index=False)
    spectra, axis = mod.load_opus_files(base['spectra_dir'], cache_path=str(output / 'opus_spectra_cache.joblib'), n_jobs=4)
    all_reports = []; searches = []; payload = {}; settings = {}
    for prop in properties:
        cfg = mod.build_property_config(base, prop, metadata)
        if cfg['metadata_warning']:
            raise ValueError(f'{prop}: {cfg["metadata_warning"]}')
        frame = original_frames[prop]
        if not previous_audit.empty:
            frame = frame[~frame['File Name'].isin(previous_audit.loc[previous_audit.Property.eq(prop), 'File Name'])]
        report, search, data, _, _ = analyse_property(frame, spectra, axis, cfg, (bounds or {}).get(prop))
        report.insert(0, 'Property', prop); search['Property'] = prop
        all_reports.append(report); searches.append(search); payload[prop] = data
        data['basis'] = f'PCA fitted before exclusions in pass {iteration}'
        settings[prop] = {'transform': cfg['transform'], 'units': cfg['units'], 'exclude_co2': cfg['exclude_co2'],
                          'bounds': (bounds or {}).get(prop, {}), 'initial_samples': original_frames[prop].Sample.nunique()}
    review = pd.concat(all_reports, ignore_index=True)
    if not previous_audit.empty:
        locked = previous_audit[COLS + ['Property']].copy()
        locked['Confirm'] = 'YES'; locked['Reason'] = previous_audit.Reasons
        locked['Unit'] = np.where(locked.Reason.eq('concentration'), 'sample', 'spectrum')
        if locked.Reason.str.contains(',').any():
            raise ValueError('Multiple-reason exclusions require a fresh cumulative review rather than a new PCA pass.')
        review = pd.concat([review, locked], ignore_index=True)
    with pd.ExcelWriter(output / 'review.xlsx', engine='openpyxl') as writer:
        review.to_excel(writer, sheet_name='Review', index=False)
        pd.concat(searches, ignore_index=True).to_excel(writer, sheet_name='Preprocessing Search', index=False)
        historical_comparison(source, properties).to_excel(writer, sheet_name='Historical Comparison', index=False)
        pd.DataFrame({'Instructions': ['Set Confirm to YES to exclude, NO to retain, or leave blank.',
            'Reason: range, redundancy, spectral_quality, concentration. Historical replay requires its explicit switch.',
            'Unit: spectrum or sample. Concentration requires sample. All diagnostics are review candidates only.',
            'Concentration cap is combined across both halves and all review passes.',
            'Nearest neighbours are ranked similarities, not automatic bad-spectrum flags.']}).to_excel(writer, sheet_name='Read Me', index=False)
        ws = writer.sheets['Review']; ws.freeze_panes = 'E2'; ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(42, max(16, len(str(col[0].value))+2))
    write_json(output / 'pca_data.json', payload)
    write_json(output / 'pca_baseline.json', payload)
    if previous_payload is not None:
        write_json(output / 'previous_pass_pca_data.json', previous_payload)
        render_viewer(previous_payload, output / 'previous_pass_PCA.html')
    render_viewer(payload, output / 'PCA_viewer.html'); export_figures(payload, output)
    files = sorted(set(f for p in properties for f in original_frames[p]['File Name']))
    spectral_paths = {p.name: p for p in Path(base['spectra_dir']).rglob('*') if p.is_file() and p.name in files}
    write_json(output / 'manifest.json', {'version': 1, 'source': str(source_path), 'source_sha256': fingerprint(source_path),
        'properties': properties, 'skipped_properties': selection_issues, 'pass': iteration, 'settings': settings,
        'prior_exclusions': previous_audit.fillna('').to_dict(orient='records'),
        'diagnostic_settings': {'sg_window': 11, 'sg_polyorder': 2, 'pca_max_rank': 15, 'diagnostic_max_rank': 15, 'residual_threshold': 2.5, 'concentration_sample_fraction': float(concentration_limit_percent_setting() / Decimal(100))},
        'spectra_sha256': {f: fingerprint(spectral_paths[f]) for f in files},
        'method': 'SOP-inspired; Python thresholds; manual confirmation required',
        'module_sha256': fingerprint(Path(__file__)), 'backend_sha256': fingerprint(Path(mod.__file__))})
    print(f'Review and PCA viewer saved in {output}. No records were excluded.')


def apply_review(review_dir: Path, destination: Path, historical_replay: bool = False,
                 concentration_limit: str | None = None) -> None:
    limit = concentration_limit_setting(concentration_limit)
    manifest = json.loads((review_dir / 'manifest.json').read_text(encoding='utf-8'))
    source_path = Path(manifest['source'])
    if fingerprint(source_path) != manifest['source_sha256']:
        raise ValueError('Source workbook changed since analysis; analyse again.')
    if destination.resolve() == source_path.resolve() or destination.exists():
        raise ValueError('Choose a new output workbook; existing workbooks are not overwritten.')
    source = pd.read_excel(source_path, sheet_name=None)
    review = pd.read_excel(review_dir / 'review.xlsx', sheet_name='Review')
    if not set(review.Property).issubset(manifest['properties']):
        raise ValueError('Unknown property in review.')
    confirmed_rows(review)
    audits = []; retained = {}
    for prop in manifest['properties']:
        retained[prop], audit = apply_confirmations(source[prop], review, prop, historical_replay, limit)
        audits.append(audit)
    audit = pd.concat(audits, ignore_index=True)
    for prior in manifest.get('prior_exclusions', []):
        matches = audit[(audit.Property == prior['Property']) & (audit['File Name'] == prior['File Name'])]
        if len(matches) != 1 or matches.iloc[0]['Reasons'] != prior['Reasons']:
            raise ValueError('Earlier pass decisions are locked; start a fresh analysis to revise them.')
    excluded = {p: set(audit.loc[audit.Property.eq(p), 'File Name']) for p in manifest['properties']}
    import openpyxl
    book = openpyxl.load_workbook(source_path)
    for prop in manifest['properties']:
        ws = book[prop]; headers = [c.value for c in ws[1]]; idx = headers.index('File Name') + 1
        for r in range(ws.max_row, 1, -1):
            if ws.cell(r, idx).value in excluded[prop]:
                ws.delete_rows(r)
    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)
    audit.to_csv(review_dir / 'confirmed_audit.csv', index=False)
    historical_comparison(source, manifest['properties'], excluded).to_csv(review_dir / 'historical_comparison.csv', index=False)
    payload = json.loads((review_dir / 'pca_baseline.json').read_text(encoding='utf-8'))
    for prop, data in payload.items():
        for row in data['records']:
            row['excluded'] = row['file'] in excluded[prop]
            if row['excluded']:
                row['status'] = 'confirmed excluded'
    write_json(review_dir / 'pca_data.json', payload)
    render_viewer(payload, review_dir / 'PCA_viewer.html'); export_figures(payload, review_dir)
    provenance = dict(manifest, concentration_limit=limit, concentration_limit_percent=float(concentration_limit_percent_setting()), cleaned_workbook=str(destination.resolve()), cleaned_sha256=fingerprint(destination),
                      review_sha256=fingerprint(review_dir / 'review.xlsx'), historical_replay=historical_replay,
                      confirmed_exclusions=audit.fillna('').to_dict(orient='records'),
                      evaluation='Validation describes the precleaned dataset, not independent validation of cleanup.')
    write_json(destination.with_suffix('.cleanup.json'), provenance)
    print(f'Exported {destination}; {len(audit)} confirmed spectrum exclusions.')


def configure_modelling(cfg: dict) -> dict:
    """Integration hook; disabled mode returns the existing configuration unchanged."""
    cfg = cfg.copy()
    if not cfg.get('cleanup_enabled', False):
        return cfg
    workbook = cfg.get('cleanup_workbook')
    review_dir = cfg.get('cleanup_review_dir')
    if bool(workbook) == bool(review_dir):
        raise ValueError('Set exactly one cleanup_workbook or cleanup_review_dir.')
    if workbook:
        path = Path(workbook)
        provenance = json.loads(path.with_suffix('.cleanup.json').read_text(encoding='utf-8'))
        if fingerprint(path) != provenance['cleaned_sha256']:
            raise ValueError('Cleaned workbook no longer matches its provenance.')
        cfg['reference_excel'] = str(path)
    else:
        provenance = json.loads((Path(review_dir) / 'manifest.json').read_text(encoding='utf-8'))
        if fingerprint(provenance['source']) != provenance['source_sha256']:
            raise ValueError('Review source changed.')
        cfg['reference_excel'] = provenance['source']
    cfg['property_sheets'] = provenance['properties']; cfg['run_all_properties'] = True
    cfg['cleanup_provenance'] = provenance
    cfg['outlier_max_pct'] = float(cfg.get('cleanup_training_outlier_max_pct', 0.0))
    cfg['ref_min'] = cfg['ref_max'] = None
    cfg['output_base_dir'] = str(Path(cfg['output_base_dir']).with_name(Path(cfg['output_base_dir']).name + ' with cleanup'))
    cfg['spectra_cache'] = str(Path(cfg['output_base_dir']) / 'opus_spectra_cache_v2.joblib')
    print('Using precleaned data. Validation describes this selected dataset, not the entire cleanup procedure.')
    return cfg


def modelling_reference_frames(frames: dict, cfg: dict) -> dict:
    if not cfg.get('cleanup_enabled'):
        return frames
    if cfg.get('cleanup_review_dir'):
        review_path = Path(cfg['cleanup_review_dir']) / 'review.xlsx'
        review = pd.read_excel(review_path, sheet_name='Review')
        if not set(review.Property).issubset(frames):
            raise ValueError('Unknown property in review.')
        limit = concentration_limit_setting(cfg.get('cleanup_concentration_limit'))
        filtered = {}; audits = []
        for p, frame in frames.items():
            filtered[p], audit = apply_confirmations(frame, review, p, concentration_limit=limit)
            audits.extend(audit.fillna('').to_dict(orient='records'))
        frames = filtered
        cfg['cleanup_provenance'] = dict(cfg['cleanup_provenance'], concentration_limit=limit, concentration_limit_percent=float(concentration_limit_percent_setting()), review_sha256=fingerprint(review_path), confirmed_exclusions=audits)
    out = Path(cfg['output_base_dir']); out.mkdir(parents=True, exist_ok=True)
    write_json(out / 'cleanup_provenance.json', cfg['cleanup_provenance'])
    return frames


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    a = commands.add_parser('analyse'); a.add_argument('--source', type=Path, default=DEFAULT_SOURCE)
    a.add_argument('--output', type=Path, default=DEFAULT_OUTPUT / datetime.now().strftime('review_%Y%m%d_%H%M%S_%f'))
    selection = a.add_mutually_exclusive_group()
    selection.add_argument('--properties', nargs='+', default=None, help='Exact sheet names, or all; overrides CONFIG.')
    selection.add_argument('--all-properties', action='store_true', help='Analyse every property sheet, including OPUS subsets.')
    a.add_argument('--bounds', type=Path, help='JSON: {"202_STN": {"min": ..., "max": ...}}')
    a.add_argument('--prior', type=Path, help='Applied previous pass directory; maximum three passes')
    b = commands.add_parser('apply'); b.add_argument('--review-dir', type=Path, required=True)
    b.add_argument('--output', type=Path, required=True); b.add_argument('--historical-replay', action='store_true')
    b.add_argument('--concentration-limit', choices=['enable', 'disable'], default=None,
                   help='Enable or disable the concentration ceiling configured by concentration_limit_percent; defaults to CONFIG (disable).')
    args = parser.parse_args(sys.argv[1:] or ['analyse'])
    with threadpool_limits(limits=1):
        if args.command == 'analyse':
            analyse(args.source, args.output, 'all' if args.all_properties else args.properties, json.loads(args.bounds.read_text(encoding='utf-8')) if args.bounds else None, args.prior)
        else:
            apply_review(args.review_dir, args.output, args.historical_replay, args.concentration_limit)

if __name__ == '__main__':
    main()
